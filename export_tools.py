"""Utility helpers to export deck sheets or preset libraries."""

from __future__ import annotations

import json
from typing import Dict

from docx import Document
from reportlab.pdfgen import canvas


def export_deck_pdf(filename: str, combos: Dict[str, Dict[str, str]]) -> None:
    """Create a simple one-page PDF deck sheet."""
    c = canvas.Canvas(filename)
    y = 800
    c.setFont("Helvetica", 14)
    c.drawString(40, y, "Beyblade X Deck Sheet")
    c.setFont("Helvetica", 12)
    for name, combo in combos.items():
        y -= 20
        text = f"{name}: {combo['blade']} / {combo['ratchet']} / {combo['bit']}"
        c.drawString(40, y, text)
    c.save()


def export_deck_docx(filename: str, combos: Dict[str, Dict[str, str]]) -> None:
    doc = Document()
    doc.add_heading("Beyblade X Deck Sheet", level=1)
    for name, combo in combos.items():
        doc.add_paragraph(f"{name}: {combo['blade']} / {combo['ratchet']} / {combo['bit']}")
    doc.save(filename)


def export_presets_json(filename: str, parts: Dict[str, Dict]) -> None:
    with open(filename, "w", encoding="utf-8") as f:
        json.dump(parts, f, indent=2)

